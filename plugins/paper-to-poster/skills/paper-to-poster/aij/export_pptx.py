#!/usr/bin/env python3
"""Export a finished AIJ poster (HTML) to an editable .pptx (docs/adr/0012).

The HTML poster is the single source of truth; this is a ONE-WAY export, run
after the gates pass. It renders the poster in print-emulated Chromium exactly
like `measure` / `render_preview`, reads the geometry of everything on the
sheet, and re-expresses it inside the organisers' own template
(aij/assets/aij_template.pptx — their master, background and AIJ mark, with the
embedded fonts stripped, docs/adr/0011):

  * text (title, № , authors, section headings, paragraphs, lists, captions,
    footer)  -> native, editable text frames at the rendered geometry, typed
    by name in SB Sans Display / SB Sans Display Light (not embedded — the
    machine opening the file needs the fonts installed);
  * tables   -> native PowerPoint tables;
  * figures, logos, QR codes, inline <svg>/<canvas>, CSS background images
    -> pictures;
  * math (inline, display, and inside table cells) -> native Office
    equations: MathJax's own MathML -> OMML (mathml2omml) inside <a14:m>,
    each math-bearing shape (text frame, equation, table) wrapped in
    <mc:AlternateContent> with a rendered-PNG fallback for non-Office
    viewers. A formula that fails to convert (converter error, or a TeX
    error MathJax rendered as <merror>) turns its whole shape into that
    picture and is listed in the export report.
  * content with no native form — text mixed with block content or with an
    inline graphic — is kept as the rendered picture and reported.

Nothing on the sheet is dropped silently: every formula is counted either as
a native equation or as a picture, and every picture fallback is reported.
The organisers' instruction text boxes, arrows, sample logos and sample QR
codes are removed; only the background and the AIJ mark are kept, and the
package thumbnail is replaced by this poster's.

Needs the `[pptx]` extra:  pip install python-pptx mathml2omml lxml
(and Playwright/Chromium, like the rest of the skill).

Usage:
  python aij/export_pptx.py poster/poster.html                 # -> poster/poster.pptx
  python aij/export_pptx.py poster/poster.html -o out.pptx --report report.json
"""
from __future__ import annotations

import argparse
import copy
import io
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
sys.path.insert(0, str(SKILL_DIR / "tools"))

from _posterly import canvas as _canvas  # noqa: E402
from _posterly import render as _render  # noqa: E402

TEMPLATE = HERE / "assets" / "aij_template.pptx"

#: The AIJ canvas = the organisers' slide (6858000 x 9907588 EMU).
AIJ_CANVAS_IN = (6858000 / 914400, 9907588 / 914400)
#: Template shapes that ARE the frame and stay; everything else is removed.
KEEP_SHAPE_NAMES = ("Рисунок 10", "Рисунок 8")   # background, AIJ mark

#: Font names as the organisers' desktop fonts expose them (the template itself
#: uses the "SB Sans Display Semibold" family-per-weight convention).
TYPEFACE_LIGHT = "SB Sans Display Light"
TYPEFACE_REGULAR = "SB Sans Display"
MATH_TYPEFACE = "Cambria Math"

EMU_PER_PX = 9525            # 914400 EMU/in / 96 px/in
PT_PER_PX = 0.75
#: Extra frame width so PowerPoint's slightly different line-breaking does not
#: push a word onto an extra line (left-aligned text grows right, centred text
#: grows both ways, right-aligned grows left).
WIDTH_SLACK = 0.015
FALLBACK_SCALE = 4           # device-scale factor for fallback / SVG rasters

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # "No Style, No Grid"


def _q(tag: str) -> str:
    pfx, local = tag.split(":")
    return f"{{{NS[pfx]}}}{local}"


# --------------------------------------------------------------------------
# In-page extraction
# --------------------------------------------------------------------------

EXTRACT_JS = r"""
() => {
  const poster = document.querySelector('[data-measure-role="poster"]');
  if (!poster) return {error: 'no [data-measure-role="poster"] element'};
  const P = poster.getBoundingClientRect();
  const out = {poster: {x: P.x, y: P.y, w: P.width, h: P.height},
               lang: document.documentElement.lang || 'en',
               rects: [], images: [], snapshots: [], texts: [], tables: [], equations: [], warnings: []};
  let nextId = 0;
  const tag = (el) => { const id = String(nextId++); el.setAttribute('data-aij-export-id', id); return id; };
  const cs = (el) => getComputedStyle(el);
  const px = (v) => parseFloat(v) || 0;
  const rgba = (s) => {
    const m = s && s.match(/rgba?\(([^)]+)\)/);
    if (!m) return null;
    const p = m[1].split(',').map(t => parseFloat(t));
    const a = p.length > 3 ? p[3] : 1;
    if (a <= 0.01) return null;
    return '#' + p.slice(0, 3).map(v => Math.round(v).toString(16).padStart(2, '0')).join('').toUpperCase();
  };
  const box = (r) => ({x: r.x - P.x, y: r.y - P.y, w: r.width, h: r.height});
  const contentBox = (el) => {
    const r = el.getBoundingClientRect(), s = cs(el);
    const l = px(s.borderLeftWidth) + px(s.paddingLeft), t = px(s.borderTopWidth) + px(s.paddingTop);
    const rr = px(s.borderRightWidth) + px(s.paddingRight), b = px(s.borderBottomWidth) + px(s.paddingBottom);
    return {x: r.x + l - P.x, y: r.y + t - P.y, w: r.width - l - rr, h: r.height - t - b};
  };
  const hidden = (el) => { const s = cs(el); return s.display === 'none' || s.visibility === 'hidden' || parseFloat(s.opacity) === 0; };
  const isInline = (el) => { const d = cs(el).display; return d.startsWith('inline') || d === 'contents'; };
  const isMath = (el) => el.tagName === 'MJX-CONTAINER';
  const isDisplayMath = (el) => isMath(el) && el.getAttribute('display') === 'true';
  // Replaced / drawn content that becomes a picture (an HTML-namespace <svg>'s tagName is lowercase).
  const GRAPHIC = new Set(['IMG', 'svg', 'CANVAS', 'VIDEO', 'OBJECT', 'EMBED', 'IFRAME']);
  const isGraphic = (el) => GRAPHIC.has(el.tagName);
  // Inside a formula or an inline SVG: not HTML layout, never inspected on its own.
  const opaque = (d) => (d.closest('mjx-container') && d.tagName !== 'MJX-CONTAINER') ||
                        (d.parentElement && d.parentElement.closest('svg'));
  const SKIP = (el) => el.classList.contains('aij-bg') || el.classList.contains('aij-mark') ||
                        ['SCRIPT', 'STYLE', 'NOSCRIPT', 'TEMPLATE'].includes(el.tagName);

  // MathJax items -> MathML, keyed by their container element.
  const mathOf = new Map();
  if (window.MathJax && MathJax.startup && MathJax.startup.document) {
    for (const it of MathJax.startup.document.getMathItemsWithin(document.body)) {
      if (it.typesetRoot) mathOf.set(it.typesetRoot, {tex: it.math, mml: MathJax.startup.toMML(it.root), display: !!it.display});
    }
  }

  const lineHeightPx = (s) => {
    const lh = s.lineHeight;
    return lh === 'normal' ? 1.2 * px(s.fontSize) : px(lh);
  };
  const runStyle = (el) => {
    const s = cs(el);
    let baseline = 0, sizeEl = el;
    for (let e = el; e && e !== document.body; e = e.parentElement) {
      const va = cs(e).verticalAlign;
      if (va === 'super' || e.tagName === 'SUP') { baseline = 30000; sizeEl = e.parentElement; break; }
      if (va === 'sub' || e.tagName === 'SUB') { baseline = -25000; sizeEl = e.parentElement; break; }
      if (!isInline(e)) break;
    }
    const fam = s.fontFamily.split(',')[0].replace(/["']/g, '').trim();
    return {
      family: fam, fontPx: px(cs(sizeEl).fontSize), weight: parseInt(s.fontWeight, 10) || 400,
      italic: s.fontStyle === 'italic' || s.fontStyle === 'oblique', color: rgba(s.color) || '#000000',
      baseline, caps: s.textTransform === 'uppercase', smallcaps: s.fontVariantCaps === 'small-caps',
      spacingPx: s.letterSpacing === 'normal' ? 0 : px(s.letterSpacing),
      underline: (s.textDecorationLine || '').includes('underline'),
    };
  };
  // Inline content of a block -> runs (text / math / br), CSS white-space collapsed.
  const runsOf = (block) => {
    const runs = [];
    const rec = (node) => {
      for (const ch of node.childNodes) {
        if (ch.nodeType === 3) {
          if (ch.nodeValue.length) runs.push(Object.assign({type: 'text', text: ch.nodeValue}, runStyle(ch.parentElement)));
        } else if (ch.nodeType === 1) {
          if (hidden(ch)) continue;
          if (isMath(ch)) {
            const m = mathOf.get(ch);
            const st = runStyle(ch);
            runs.push({type: 'math', id: tag(ch), mml: m ? m.mml : null, tex: m ? m.tex : ch.textContent,
                       fontPx: st.fontPx, color: st.color});
            continue;
          }
          if (ch.tagName === 'BR') { runs.push({type: 'br'}); continue; }
          if (ch.tagName === 'IMG') { out.warnings.push('inline <img> inside text is not exported: ' + (ch.getAttribute('src') || '')); continue; }
          rec(ch);
        }
      }
    };
    rec(block);
    const ws = cs(block).whiteSpace;
    if (!ws.startsWith('pre')) {
      for (const r of runs) if (r.type === 'text') r.text = r.text.replace(/[ \t\n\r\f]+/g, ' ');
      // drop a space at the start / end of a line and double spaces across runs
      let prevSpace = true;
      for (const r of runs) {
        if (r.type === 'br') { prevSpace = true; continue; }
        if (r.type === 'math') { prevSpace = false; continue; }
        if (prevSpace) r.text = r.text.replace(/^ /, '');
        if (r.text.length) prevSpace = r.text.endsWith(' ');
      }
      for (let i = runs.length - 1; i >= 0; i--) {
        const r = runs[i];
        if (r.type === 'br') continue;
        if (r.type === 'text') { r.text = r.text.replace(/ $/, ''); if (r.text.length) break; } else break;
      }
    }
    return runs.filter(r => r.type !== 'text' || r.text.length);
  };
  const paraOf = (el, extra) => {
    const s = cs(el);
    let align = s.textAlign;
    if (align === 'start' || align === '-webkit-auto') align = 'left';
    if (align === 'end') align = 'right';
    return Object.assign({align, lineHeightPx: lineHeightPx(s), fontPx: px(s.fontSize), runs: runsOf(el)}, extra || {});
  };
  // A text block: a non-inline element whose visible content is inline text/math only
  // (no block children, no graphics — those make it a picture, see `mixed` below).
  const inlineOnly = (el) => {
    for (const d of el.querySelectorAll('*')) {
      if (opaque(d) || hidden(d)) continue;
      if (isDisplayMath(d)) return false;
      if (isGraphic(d) || d.tagName === 'TABLE') return false;
      if (!isInline(d) && !isMath(d)) return false;
    }
    return true;
  };
  const hasContent = (el) => el.textContent.trim().length > 0 || el.querySelector('mjx-container');
  // Text or inline math sitting directly next to block content / graphics — it
  // cannot become one native text frame, so the element becomes a picture.
  const mixed = (el) => [...el.childNodes].some(ch =>
      (ch.nodeType === 3 && ch.nodeValue.trim()) ||
      (ch.nodeType === 1 && !hidden(ch) && !isGraphic(ch) && isInline(ch) && hasContent(ch)) ||
      (ch.nodeType === 1 && isMath(ch) && !isDisplayMath(ch)));
  const texOf = (root) => [...root.querySelectorAll('mjx-container')].map(m => (mathOf.get(m) || {}).tex || m.textContent);

  const decor = (el) => {
    if (el === poster) return;
    const s = cs(el);
    const r = el.getBoundingClientRect();
    if (r.width < 0.5 || r.height < 0.5) return;
    const fill = rgba(s.backgroundColor);
    if (fill) out.rects.push(Object.assign(box(r), {fill, radius: px(s.borderTopLeftRadius)}));
    // CSS background images / gradients: rasterised in Python with the element's
    // content hidden, so the picture carries only the background.
    if (s.backgroundImage && s.backgroundImage !== 'none')
      out.rects.push(Object.assign(box(r), {bgimage: true, id: tag(el)}));
    const sides = [['Top', 0, 0, r.width, px(s.borderTopWidth)], ['Bottom', 0, r.height - px(s.borderBottomWidth), r.width, px(s.borderBottomWidth)],
                   ['Left', 0, 0, px(s.borderLeftWidth), r.height], ['Right', r.width - px(s.borderRightWidth), 0, px(s.borderRightWidth), r.height]];
    for (const [side, dx, dy, w, h] of sides) {
      if (s['border' + side + 'Style'] === 'none' || w <= 0 || h <= 0) continue;
      const c = rgba(s['border' + side + 'Color']);
      if (c) out.rects.push({x: r.x - P.x + dx, y: r.y - P.y + dy, w, h, fill: c, radius: 0});
    }
  };

  const tableOf = (t) => {
    const T = t.getBoundingClientRect();
    const rows = [];
    const lefts = new Set(), tops = new Set();
    for (const tr of t.rows) {
      const R = tr.getBoundingClientRect();
      const cells = [];
      for (const td of tr.cells) {
        const C = td.getBoundingClientRect(), s = cs(td);
        lefts.add(Math.round((C.x - T.x) * 10) / 10);
        cells.push({x: C.x - P.x, y: C.y - P.y, w: C.width, h: C.height, colspan: td.colSpan, rowspan: td.rowSpan,
          fill: rgba(s.backgroundColor) || rgba(cs(tr).backgroundColor),
          valign: s.verticalAlign,
          pad: {l: px(s.paddingLeft), r: px(s.paddingRight), t: px(s.paddingTop), b: px(s.paddingBottom)},
          borders: Object.fromEntries(['Top', 'Bottom', 'Left', 'Right'].map(k => [k[0].toLowerCase(),
            (s['border' + k + 'Style'] !== 'none' && px(s['border' + k + 'Width']) > 0)
              ? {w: px(s['border' + k + 'Width']), color: rgba(s['border' + k + 'Color'])} : null])),
          para: paraOf(td)});
      }
      tops.add(Math.round((R.y - T.y) * 10) / 10);
      rows.push({y: R.y - P.y, h: R.height, cells});
    }
    return Object.assign(box(T), {id: tag(t), rows, colLefts: [...lefts].sort((a, b) => a - b)});
  };

  const walk = (el) => {
    if (SKIP(el) || hidden(el)) return;
    // Content that has no native pptx form is kept as the rendered picture and
    // reported — never dropped silently.
    if (!isGraphic(el) && !isMath(el) && el.tagName !== 'TABLE' && !(!isInline(el) && hasContent(el) && inlineOnly(el))
        && mixed(el)) {
      const r = el.getBoundingClientRect();
      out.snapshots.push(Object.assign(box(r), {id: tag(el), tex: texOf(el),
        text: el.textContent.replace(/\s+/g, ' ').trim().slice(0, 80)}));
      out.warnings.push('text mixed with block content or an inline graphic is exported as a picture: "'
                        + el.textContent.replace(/\s+/g, ' ').trim().slice(0, 60) + '"');
      return;
    }
    decor(el);
    if (isGraphic(el)) {
      const b = contentBox(el), img = el.tagName === 'IMG';
      out.images.push(Object.assign(b, {id: tag(el), src: img ? (el.currentSrc || el.src) : '',
        natW: img ? el.naturalWidth : 0, natH: img ? el.naturalHeight : 0, fit: cs(el).objectFit}));
      return;
    }
    if (el.tagName === 'TABLE') { out.tables.push(tableOf(el)); return; }
    if (isDisplayMath(el)) {
      const m = mathOf.get(el), s = cs(el);
      out.equations.push(Object.assign(contentBox(el), {id: tag(el), mml: m ? m.mml : null, tex: m ? m.tex : el.textContent,
        fontPx: px(s.fontSize), color: rgba(s.color) || '#000000', align: 'center'}));
      return;
    }
    if (isMath(el)) {   // inline math whose parent is inline too, under a block-content parent
      const r = el.getBoundingClientRect();
      out.snapshots.push(Object.assign(box(r), {id: tag(el), tex: [(mathOf.get(el) || {}).tex || el.textContent], text: ''}));
      out.warnings.push('an inline formula outside any text block is exported as a picture');
      return;
    }
    if ((el.tagName === 'UL' || el.tagName === 'OL') && [...el.children].every(li => li.tagName === 'LI' && inlineOnly(li))) {
      const s = cs(el), items = [...el.children].filter(li => !hidden(li));
      if (!items.length) return;
      const ulBox = contentBox(el);
      const indent = px(s.paddingLeft);
      let start = el.tagName === 'OL' ? (parseInt(el.getAttribute('start') || '1', 10)) : 1;
      const paras = items.map((li, i) => {
        const ls = cs(li), prev = i ? items[i - 1] : null;
        const before = i ? px(ls.marginTop) + (prev ? px(cs(prev).marginBottom) : 0) : 0;
        return paraOf(li, {bullet: ls.listStyleType, spaceBeforePx: before, indentPx: indent, startAt: start});
      });
      out.texts.push(Object.assign({x: ulBox.x - indent, y: ulBox.y, w: ulBox.w + indent, h: ulBox.h}, {id: tag(el), paragraphs: paras}));
      return;
    }
    if (!isInline(el) && hasContent(el) && inlineOnly(el)) {
      out.texts.push(Object.assign(contentBox(el), {id: tag(el), paragraphs: [paraOf(el)]}));
      return;
    }
    for (const ch of el.children) walk(ch);
  };
  walk(poster);
  return out;
}
"""


# --------------------------------------------------------------------------
# Math: MathML -> OMML
# --------------------------------------------------------------------------

def _default_mathml_to_omml(mml: str):
    """MathJax MathML -> an <m:oMath> lxml element (raises on failure)."""
    import mathml2omml
    from lxml import etree

    mml = re.sub(r">\s+<", "><", mml.strip())
    omml = mathml2omml.convert(mml)
    root = etree.fromstring(f'<root xmlns:m="{NS["m"]}">{omml}</root>')
    om = root.find(_q("m:oMath"))
    if om is None or not len(om):
        raise ValueError("converter produced no <m:oMath> content")
    normalize_omml(om)
    return om


def normalize_omml(om) -> None:
    """Repair known mathml2omml gaps so the OMML is schema-valid.

    * <m:rad> without <m:deg> (plain \\sqrt): ECMA-376 requires <m:deg>; add
      an empty one and hide it (radPr/degHide), as Office itself writes it.
    """
    from lxml import etree

    for rad in om.iter(_q("m:rad")):
        if rad.find(_q("m:deg")) is None:
            rad_pr = rad.find(_q("m:radPr"))
            if rad_pr is None:
                rad_pr = etree.Element(_q("m:radPr"))
                rad.insert(0, rad_pr)
            if rad_pr.find(_q("m:degHide")) is None:
                dh = etree.SubElement(rad_pr, _q("m:degHide"))
                dh.set(_q("m:val"), "1")
            rad.insert(list(rad).index(rad_pr) + 1, etree.Element(_q("m:deg")))


def _style_math_runs(om, size_pt: float, color: str) -> None:
    """Give every <m:r> a DrawingML <a:rPr> (Cambria Math, size, colour), the
    way PowerPoint writes math runs. Inserted right after <m:rPr>."""
    from lxml import etree

    for r in om.iter(_q("m:r")):
        sty = r.find(f"{_q('m:rPr')}/{_q('m:sty')}")
        val = sty.get(_q("m:val")) if sty is not None else None
        rpr = etree.Element(_q("a:rPr"), lang="en-US", sz=str(int(round(size_pt * 100))))
        if val in ("i", "bi"):
            rpr.set("i", "1")
        if val in ("b", "bi"):
            rpr.set("b", "1")
        fill = etree.SubElement(rpr, _q("a:solidFill"))
        etree.SubElement(fill, _q("a:srgbClr"), val=color.lstrip("#"))
        etree.SubElement(rpr, _q("a:latin"), typeface=MATH_TYPEFACE, panose="02040503050406030204",
                         pitchFamily="18", charset="0")
        etree.SubElement(rpr, _q("a:cs"), typeface=MATH_TYPEFACE, panose="02040503050406030204",
                         pitchFamily="18", charset="0")
        m_rpr = r.find(_q("m:rPr"))
        r.insert(list(r).index(m_rpr) + 1 if m_rpr is not None else 0, rpr)


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

@dataclass
class ExportReport:
    pptx: str = ""
    text_frames: int = 0
    tables: int = 0
    pictures: int = 0
    shapes_decor: int = 0
    equations_native: int = 0
    equations_inline_native: int = 0
    equations_as_picture: list[dict] = field(default_factory=list)
    frames_as_picture: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return dict(self.__dict__)


def _emu(v_px: float) -> int:
    return int(round(v_px * EMU_PER_PX))


def _typeface(run: dict) -> tuple[str, bool]:
    """(typeface, bold) for a run, mapping SB Sans Display weights to the
    organisers' font names."""
    fam = run.get("family") or ""
    w = run.get("weight", 400)
    if "sb sans" in fam.lower() or not fam:
        if w <= 350:
            return TYPEFACE_LIGHT, False
        return TYPEFACE_REGULAR, w >= 650
    return fam, w >= 650


def _lang(text: str, default: str) -> str:
    if re.search(r"[Ѐ-ӿ]", text):
        return "ru-RU"
    return "en-US" if default.lower().startswith("en") else default


class _Exporter:
    def __init__(self, page, data: dict, html_path: Path, mathml_to_omml: Callable[[str], Any]):
        from pptx import Presentation

        self.page = page
        self.data = data
        self.html_path = html_path
        self.to_omml = mathml_to_omml
        self.report = ExportReport()
        self.prs = Presentation(str(TEMPLATE))
        self.slide = self.prs.slides[0]
        self.lang = data.get("lang") or "en"

    # ---- template -------------------------------------------------------
    def strip_template(self) -> None:
        tree = self.slide.shapes._spTree
        kept = set()
        for shp in list(self.slide.shapes):
            if shp.name in KEEP_SHAPE_NAMES:
                kept.add(shp.name)
                continue
            tree.remove(shp._element)
        missing = set(KEEP_SHAPE_NAMES) - kept
        if missing:
            raise SystemExit(f"error: {TEMPLATE.name} lacks frame shape(s) {sorted(missing)} — template changed?")
        xml = self.slide._element.xml
        for rel in list(self.slide.part.rels.values()):
            if rel.reltype.endswith("/image") and f'"{rel.rId}"' not in xml:
                self.slide.part.drop_rel(rel.rId)

    # ---- helpers ---------------------------------------------------------
    def _shot(self, x: float, y: float, w: float, h: float) -> bytes:
        P = self.data["poster"]
        return self.page.screenshot(clip={"x": P["x"] + x, "y": P["y"] + y, "width": max(w, 1), "height": max(h, 1)},
                                    omit_background=True, type="png")

    def _add_background(self, r: dict) -> None:
        """A CSS background image / gradient: rasterise the element with its
        own content hidden, so the picture is only the background."""
        sel = f'[data-aij-export-id="{r["id"]}"]'
        css = f"{sel} * {{ visibility: hidden !important; }} {sel} {{ color: transparent !important; }}"
        self.page.evaluate("(css) => { const s = document.createElement('style'); s.id = 'aij-bg-only';"
                           " s.textContent = css; document.head.appendChild(s); }", css)
        try:
            png = self._shot(r["x"], r["y"], r["w"], r["h"])
        finally:
            self.page.evaluate("() => { const s = document.getElementById('aij-bg-only'); if (s) s.remove(); }")
        self.slide.shapes.add_picture(io.BytesIO(png), _emu(r["x"]), _emu(r["y"]), _emu(r["w"]), _emu(r["h"]))
        self.report.pictures += 1

    def _add_snapshot(self, s: dict) -> None:
        """Content with no native pptx form (text mixed with block content or an
        inline graphic): the rendered picture, reported."""
        self.slide.shapes.add_picture(io.BytesIO(self._shot(s["x"], s["y"], s["w"], s["h"])),
                                      _emu(s["x"]), _emu(s["y"]), _emu(s["w"]), _emu(s["h"]))
        self.report.pictures += 1
        self.report.frames_as_picture.append({"reason": "text mixed with block content or an inline graphic",
                                              "text": s.get("text", "")})
        for tex in s.get("tex", []):
            self.report.equations_as_picture.append(
                {"tex": tex, "error": "inside content exported as a picture (mixed inline/block content)"})

    def _add_rect(self, r: dict) -> None:
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        if r.get("bgimage"):
            self._add_background(r)
            return

        rounded = r.get("radius", 0) > 0.3
        shp = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                          _emu(r["x"]), _emu(r["y"]), max(_emu(r["w"]), 1), max(_emu(r["h"]), 1))
        if rounded:
            shp.adjustments[0] = min(0.5, r["radius"] / max(1e-6, min(r["w"], r["h"])))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(r["fill"].lstrip("#"))
        shp.line.fill.background()
        shp.shadow.inherit = False
        # no text / style effects on a decoration shape
        style = shp._element.find(_q("p:style"))
        if style is not None:
            shp._element.remove(style)
        self.report.shapes_decor += 1

    def _add_image(self, im: dict) -> None:
        src = im["src"]
        stream = None
        path = None
        if src.startswith("file://"):
            from urllib.parse import unquote, urlparse
            path = Path(unquote(urlparse(src).path))
        natural_ar = (im["natW"] / im["natH"]) if im.get("natW") and im.get("natH") else None
        box_ar = im["w"] / im["h"] if im["h"] else None
        raster_ok = path is not None and path.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp")
        if raster_ok and natural_ar and box_ar and abs(natural_ar / box_ar - 1) < 0.01:
            stream = str(path)
        else:
            # SVG, data: URI, remote, or object-fit cropping: rasterise what Chromium drew.
            stream = io.BytesIO(self._shot(im["x"], im["y"], im["w"], im["h"]))
        self.slide.shapes.add_picture(stream, _emu(im["x"]), _emu(im["y"]), _emu(im["w"]), _emu(im["h"]))
        self.report.pictures += 1

    # ---- text ------------------------------------------------------------
    def _fill_paragraph(self, p_el, para: dict, math_ok: dict[str, Any]) -> None:
        from lxml import etree

        a = lambda t: _q("a:" + t)  # noqa: E731
        ppr = p_el.find(a("pPr"))
        if ppr is None:
            ppr = etree.Element(a("pPr"))
            p_el.insert(0, ppr)
        for ch in list(ppr):
            ppr.remove(ch)
        algn = {"left": "l", "center": "ctr", "right": "r", "justify": "just"}.get(para.get("align"), "l")
        ppr.set("algn", algn)
        bullet = para.get("bullet")
        if bullet and bullet != "none":
            ind = _emu(para.get("indentPx", 0))
            ppr.set("marL", str(ind))
            ppr.set("indent", str(-ind))
        lh = etree.SubElement(ppr, a("lnSpc"))
        etree.SubElement(lh, a("spcPts"), val=str(int(round(para["lineHeightPx"] * PT_PER_PX * 100))))
        sb = etree.SubElement(ppr, a("spcBef"))
        etree.SubElement(sb, a("spcPts"), val=str(int(round(para.get("spaceBeforePx", 0) * PT_PER_PX * 100))))
        sa = etree.SubElement(ppr, a("spcAft"))
        etree.SubElement(sa, a("spcPts"), val="0")
        if bullet and bullet != "none":
            if bullet in ("decimal", "decimal-leading-zero"):
                etree.SubElement(ppr, a("buFont"), typeface="+mj-lt")
                etree.SubElement(ppr, a("buAutoNum"), type="arabicPeriod", startAt=str(para.get("startAt", 1)))
            else:
                char = {"circle": "◦", "square": "▪"}.get(bullet, "•")
                etree.SubElement(ppr, a("buFont"), typeface="Arial")
                etree.SubElement(ppr, a("buChar"), char=char)
        else:
            etree.SubElement(ppr, a("buNone"))
        for r in list(p_el):
            if r is not ppr:
                p_el.remove(r)

        last_size = para.get("fontPx", 9.33) * PT_PER_PX
        for run in para["runs"]:
            if run["type"] == "br":
                br = etree.SubElement(p_el, a("br"))
                etree.SubElement(br, a("rPr"), lang=_lang("", self.lang), sz=str(int(round(last_size * 100))))
                continue
            if run["type"] == "math":
                om = math_ok[run["id"]]
                om = copy.deepcopy(om)
                _style_math_runs(om, run["fontPx"] * PT_PER_PX, run["color"])
                m_wrap = etree.SubElement(p_el, _q("a14:m"), nsmap={"a14": NS["a14"]})
                m_wrap.append(om)
                self.report.equations_inline_native += 1
                continue
            last_size = run["fontPx"] * PT_PER_PX
            face, bold = _typeface(run)
            r_el = etree.SubElement(p_el, a("r"))
            rpr = etree.SubElement(r_el, a("rPr"), lang=_lang(run["text"], self.lang),
                                   sz=str(int(round(last_size * 100))))
            if bold:
                rpr.set("b", "1")
            if run.get("italic"):
                rpr.set("i", "1")
            if run.get("underline"):
                rpr.set("u", "sng")
            if run.get("caps"):
                rpr.set("cap", "all")
            elif run.get("smallcaps"):
                rpr.set("cap", "small")
            if run.get("spacingPx"):
                rpr.set("spc", str(int(round(run["spacingPx"] * PT_PER_PX * 100))))
            if run.get("baseline"):
                rpr.set("baseline", str(run["baseline"]))
            fill = etree.SubElement(rpr, a("solidFill"))
            etree.SubElement(fill, a("srgbClr"), val=run["color"].lstrip("#"))
            for slot in ("latin", "ea", "cs"):
                etree.SubElement(rpr, a(slot), typeface=face)
            t = etree.SubElement(r_el, a("t"))
            t.text = run["text"]
        end = etree.SubElement(p_el, a("endParaRPr"), lang=_lang("", self.lang), sz=str(int(round(last_size * 100))))
        del end

    def _convert(self, item: dict):
        """MathML of one MathJax item -> <m:oMath> (raises on any failure)."""
        mml = item.get("mml")
        if not mml:
            raise ValueError("MathJax produced no MathML for this formula")
        # MathJax renders bad TeX instead of failing: syntax errors as <merror>,
        # undefined macros (its `noundefined` extension) as red <mtext>\name.
        m = re.search(r"<merror[^>]*>(.*?)</merror>", mml, re.DOTALL)
        if m:
            msg = re.sub(r"<[^>]+>|\s+", " ", m.group(1)).strip()
            raise ValueError(f"MathJax could not typeset this TeX ({msg})")
        m = re.search(r'<mtext[^>]*mathcolor="red"[^>]*>(\\[A-Za-z]+)</mtext>', mml)
        if m:
            raise ValueError(f"undefined TeX macro {m.group(1)}")
        return self.to_omml(mml)

    def _math_elements(self, paragraphs: list[dict]) -> tuple[dict[str, Any], list[dict], list[dict]]:
        """Convert every formula in `paragraphs`: (ok by run id, failures, all math runs)."""
        ok, failed, runs = {}, [], []
        for para in paragraphs:
            for run in para["runs"]:
                if run["type"] != "math":
                    continue
                runs.append(run)
                try:
                    ok[run["id"]] = self._convert(run)
                except Exception as e:  # noqa: BLE001 — any converter failure degrades to a picture
                    failed.append({"tex": run.get("tex"), "error": f"{type(e).__name__}: {e}"})
        return ok, failed, runs

    def _report_as_picture(self, failed: list[dict], runs: list[dict], ok: dict, where: str, text: str) -> None:
        """Every formula of a shape that fell back to a picture is accounted for:
        the failing ones with their error, the converted ones as carried along."""
        self.report.equations_as_picture.extend(failed)
        for run in runs:
            if run["id"] in ok:
                self.report.equations_as_picture.append(
                    {"tex": run.get("tex"), "error": f"converted, but kept as a picture with its {where} "
                                                     "(another formula in it failed)"})
        self.report.frames_as_picture.append({"reason": f"equation conversion failed ({where})", "text": text})

    def _textbox(self, t: dict):
        x, w = t["x"], t["w"]
        slack = w * WIDTH_SLACK
        align = t["paragraphs"][0].get("align", "left") if t["paragraphs"] else "left"
        if align == "center":
            x -= slack / 2
        elif align == "right":
            x -= slack
        w += slack
        shp = self.slide.shapes.add_textbox(_emu(x), _emu(t["y"]), _emu(w), max(_emu(t["h"]), 1))
        tf = shp.text_frame
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        return shp

    def _add_text(self, t: dict) -> None:
        from lxml import etree

        math_ok, failed, math_runs = self._math_elements(t["paragraphs"])
        has_math = bool(math_runs)
        if failed:
            # One formula could not become a native equation: keep the whole
            # frame faithful as the rendered picture and say so.
            self.slide.shapes.add_picture(io.BytesIO(self._shot(t["x"], t["y"], t["w"], t["h"])),
                                          _emu(t["x"]), _emu(t["y"]), _emu(t["w"]), _emu(t["h"]))
            self.report.pictures += 1
            self._report_as_picture(failed, math_runs, math_ok, "text frame", _plain(t)[:80])
            return
        shp = self._textbox(t)
        body = shp.text_frame._txBody
        for p in body.findall(_q("a:p")):
            body.remove(p)
        for para in t["paragraphs"]:
            p_el = etree.SubElement(body, _q("a:p"))
            self._fill_paragraph(p_el, para, math_ok)
        self.report.text_frames += 1
        if has_math:
            self._wrap_alternate(shp, t)

    def _add_equation(self, eq: dict) -> None:
        from lxml import etree

        try:
            om = self._convert(eq)
        except Exception as e:  # noqa: BLE001
            self.slide.shapes.add_picture(io.BytesIO(self._shot(eq["x"], eq["y"], eq["w"], eq["h"])),
                                          _emu(eq["x"]), _emu(eq["y"]), _emu(eq["w"]), _emu(eq["h"]))
            self.report.pictures += 1
            self.report.equations_as_picture.append({"tex": eq.get("tex"), "error": f"{type(e).__name__}: {e}"})
            return
        om = copy.deepcopy(om)
        _style_math_runs(om, eq["fontPx"] * PT_PER_PX, eq["color"])
        shp = self._textbox(dict(eq, paragraphs=[{"align": "center"}]))
        body = shp.text_frame._txBody
        for p in body.findall(_q("a:p")):
            body.remove(p)
        p_el = etree.SubElement(body, _q("a:p"))
        ppr = etree.SubElement(p_el, _q("a:pPr"), algn="ctr")
        etree.SubElement(ppr, _q("a:buNone"))
        wrap = etree.SubElement(p_el, _q("a14:m"), nsmap={"a14": NS["a14"]})
        para = etree.SubElement(wrap, _q("m:oMathPara"))
        para_pr = etree.SubElement(para, _q("m:oMathParaPr"))
        etree.SubElement(para_pr, _q("m:jc"), {_q("m:val"): "centerGroup"})
        para.append(om)
        etree.SubElement(p_el, _q("a:endParaRPr"), lang="en-US", sz=str(int(round(eq["fontPx"] * PT_PER_PX * 100))))
        self.report.equations_native += 1
        self._wrap_alternate(shp, eq)

    def _wrap_alternate(self, shp, box: dict) -> None:
        """Wrap a math-bearing <p:sp> as PowerPoint does: mc:Choice (a14) holds
        the native equation; mc:Fallback holds the same shape filled with the
        rendered picture (for viewers without Office math)."""
        from lxml import etree

        sp = shp._element
        png = self._shot(box["x"], box["y"], box["w"], box["h"])
        _part, rid = self.slide.part.get_or_add_image_part(io.BytesIO(png))
        fb = copy.deepcopy(sp)
        sppr = fb.find(_q("p:spPr"))
        xfrm = sppr.find(_q("a:xfrm"))
        # Fallback picture covers the rendered box exactly (no width slack).
        xfrm.find(_q("a:off")).set("x", str(_emu(box["x"])))
        xfrm.find(_q("a:ext")).set("cx", str(_emu(box["w"])))
        for ch in list(sppr):
            if ch.tag in (_q("a:noFill"), _q("a:solidFill"), _q("a:blipFill")):
                sppr.remove(ch)
        blip_fill = etree.Element(_q("a:blipFill"))
        etree.SubElement(blip_fill, _q("a:blip"), {_q("r:embed"): rid})
        stretch = etree.SubElement(blip_fill, _q("a:stretch"))
        etree.SubElement(stretch, _q("a:fillRect"))
        geom = sppr.find(_q("a:prstGeom"))
        sppr.insert(list(sppr).index(geom) + 1 if geom is not None else len(sppr), blip_fill)
        tx = fb.find(_q("p:txBody"))
        for p in tx.findall(_q("a:p")):
            tx.remove(p)
        etree.SubElement(tx, _q("a:p"))

        self._alternate(sp, fb)

    @staticmethod
    def _alternate(choice_el, fallback_el) -> None:
        """Replace `choice_el` in place by mc:AlternateContent: Choice (needs a14,
        i.e. Office math) = `choice_el`, Fallback = `fallback_el`."""
        from lxml import etree

        ac = etree.Element(_q("mc:AlternateContent"), nsmap={"mc": NS["mc"]})
        choice = etree.SubElement(ac, _q("mc:Choice"), Requires="a14", nsmap={"a14": NS["a14"]})
        fallback = etree.SubElement(ac, _q("mc:Fallback"))
        parent = choice_el.getparent()
        parent.replace(choice_el, ac)
        choice.append(choice_el)
        if fallback_el.getparent() is not None:
            fallback_el.getparent().remove(fallback_el)
        fallback.append(fallback_el)

    def _set_thumbnail(self) -> None:
        """The package thumbnail (file browsers) = this poster, not the template."""
        from PIL import Image

        P = self.data["poster"]
        shot = self.page.screenshot(clip={"x": P["x"], "y": P["y"], "width": P["w"], "height": P["h"]}, type="png")
        im = Image.open(io.BytesIO(shot)).convert("RGB")
        im.thumbnail((256, 256))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)
        for part in self.prs.part.package.iter_parts():
            if str(part.partname) == "/docProps/thumbnail.jpeg":
                part._blob = buf.getvalue()

    def _add_table(self, tb: dict) -> None:
        from lxml import etree
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR

        n_rows = len(tb["rows"])
        lefts = [v for v in tb["colLefts"]]
        n_cols = max(len(lefts), 1)
        if not n_rows:
            return
        cells = [c for row in tb["rows"] for c in row["cells"]]
        math_ok, failed, math_runs = self._math_elements([c["para"] for c in cells])
        if failed:
            # A formula in a cell cannot become a native equation: the whole table
            # stays the rendered picture (faithful, not editable) and is reported.
            self.slide.shapes.add_picture(io.BytesIO(self._shot(tb["x"], tb["y"], tb["w"], tb["h"])),
                                          _emu(tb["x"]), _emu(tb["y"]), _emu(tb["w"]), _emu(tb["h"]))
            self.report.pictures += 1
            first = " | ".join(_plain({"paragraphs": [c["para"]]}) for c in cells[:3])
            self._report_as_picture(failed, math_runs, math_ok, "table", first[:80])
            return
        gf = self.slide.shapes.add_table(n_rows, n_cols, _emu(tb["x"]), _emu(tb["y"]), _emu(tb["w"]), _emu(tb["h"]))
        table = gf.table
        tbl_pr = table._tbl.find(_q("a:tblPr"))
        for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
            tbl_pr.attrib.pop(attr, None)
        style_id = tbl_pr.find(_q("a:tableStyleId"))
        if style_id is None:
            style_id = etree.SubElement(tbl_pr, _q("a:tableStyleId"))
        style_id.text = NO_STYLE_TABLE
        edges = lefts + [tb["w"]]
        for i in range(n_cols):
            table.columns[i].width = max(_emu(edges[i + 1] - edges[i]), 1)
        occupied: set[tuple[int, int]] = set()
        for ri, row in enumerate(tb["rows"]):
            table.rows[ri].height = max(_emu(row["h"]), 1)
            ci = 0
            for cell in row["cells"]:
                while (ri, ci) in occupied:
                    ci += 1
                rel_x = round((cell["x"] - tb["x"]) * 10) / 10
                ci = min(range(n_cols), key=lambda k: abs(lefts[k] - rel_x)) if lefts else ci
                c = table.cell(ri, ci)
                span_r, span_c = max(cell["rowspan"], 1), max(cell["colspan"], 1)
                for dr in range(span_r):
                    for dc in range(span_c):
                        occupied.add((ri + dr, ci + dc))
                if span_r > 1 or span_c > 1:
                    c.merge(table.cell(min(ri + span_r, n_rows) - 1, min(ci + span_c, n_cols) - 1))
                pad = cell["pad"]
                c.margin_left, c.margin_right = _emu(pad["l"]), _emu(pad["r"])
                c.margin_top, c.margin_bottom = _emu(pad["t"]), _emu(pad["b"])
                c.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}.get(cell["valign"], MSO_ANCHOR.MIDDLE)
                if cell.get("fill"):
                    c.fill.solid()
                    c.fill.fore_color.rgb = RGBColor.from_string(cell["fill"].lstrip("#"))
                else:
                    c.fill.background()
                tc_pr = c._tc.get_or_add_tcPr()
                for side, tag in (("l", "a:lnL"), ("r", "a:lnR"), ("t", "a:lnT"), ("b", "a:lnB")):
                    old = tc_pr.find(_q(tag))
                    if old is not None:
                        tc_pr.remove(old)
                    bd = cell["borders"].get(side)
                    ln = etree.Element(_q(tag), w=str(_emu(bd["w"])) if bd else "0")
                    if bd and bd.get("color"):
                        sf = etree.SubElement(ln, _q("a:solidFill"))
                        etree.SubElement(sf, _q("a:srgbClr"), val=bd["color"].lstrip("#"))
                    else:
                        etree.SubElement(ln, _q("a:noFill"))
                    # line elements must precede the cell fill in CT_TableCellProperties
                    tc_pr.insert(("l", "r", "t", "b").index(side), ln)
                tx = c._tc.get_or_add_txBody()
                for p in tx.findall(_q("a:p")):
                    tx.remove(p)
                p_el = etree.SubElement(tx, _q("a:p"))
                self._fill_paragraph(p_el, cell["para"], math_ok)
        self.report.tables += 1
        if math_runs:
            # Office-math table: Choice = the native table; Fallback = the
            # rendered picture of it, for viewers without Office math.
            pic = self.slide.shapes.add_picture(io.BytesIO(self._shot(tb["x"], tb["y"], tb["w"], tb["h"])),
                                                _emu(tb["x"]), _emu(tb["y"]), _emu(tb["w"]), _emu(tb["h"]))
            self._alternate(gf._element, pic._element)

    # ---- driver ------------------------------------------------------------
    def run(self, out_path: Path) -> ExportReport:
        self.strip_template()
        d = self.data
        for r in d["rects"]:
            self._add_rect(r)
        for im in d["images"]:
            self._add_image(im)
        for sn in d.get("snapshots", []):
            self._add_snapshot(sn)
        for tb in d["tables"]:
            self._add_table(tb)
        for t in d["texts"]:
            self._add_text(t)
        for eq in d["equations"]:
            self._add_equation(eq)
        self.report.warnings.extend(d.get("warnings", []))
        self._set_thumbnail()
        self.prs.save(str(out_path))
        self.report.pptx = str(out_path)
        return self.report


def _plain(t: dict) -> str:
    text = " ".join(r.get("text", "") for p in t["paragraphs"] for r in p["runs"] if r["type"] == "text")
    return re.sub(r"\s+", " ", text).strip()


def export(html_path: Path, out_path: Path, *, mathml_to_omml: Callable[[str], Any] | None = None,
           mathjax_timeout_ms: int = 15000) -> ExportReport:
    """Render `html_path` and write the editable .pptx to `out_path`."""
    from playwright.sync_api import sync_playwright

    html_path = Path(html_path).resolve()
    canvas = _canvas.read_canvas_from_html(html_path)
    if canvas is None or abs(canvas[0] - AIJ_CANVAS_IN[0]) > 0.02 or abs(canvas[1] - AIJ_CANVAS_IN[1]) > 0.02:
        raise SystemExit(f"error: {html_path.name} is not an AIJ poster (@page must be 190.5mm 275.2mm; "
                         f"got {canvas}). The pptx export exists only for the AIJ format.")
    viewport = _canvas.viewport_for(canvas)
    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            ctx = browser.new_context(viewport={"width": viewport[0], "height": viewport[1]},
                                      device_scale_factor=FALLBACK_SCALE)
            page = ctx.new_page()
            page.emulate_media(media="print")
            page.goto(html_path.as_uri(), timeout=mathjax_timeout_ms)
            try:
                page.wait_for_load_state("networkidle", timeout=mathjax_timeout_ms)
            except Exception:  # noqa: BLE001 — settle_page reports MathJax problems below
                pass
            settle = _render.settle_page(page, mathjax_timeout_ms=mathjax_timeout_ms)
            problem = _render.hard_fail_on_settle_problems(settle, mathjax_timeout_ms=mathjax_timeout_ms)
            if problem:
                raise SystemExit(f"error: {problem}")
            fonts_ok = page.evaluate(
                "() => ['300','400','700'].every(w => document.fonts.check(w + ' 10px \"SB Sans Display\"'))")
            data = page.evaluate(EXTRACT_JS)
            if data.get("error"):
                raise SystemExit(f"error: {data['error']}")
            ex = _Exporter(page, data, html_path, mathml_to_omml or _default_mathml_to_omml)
            if not fonts_ok:
                ex.report.warnings.append("SB Sans Display did not load (CDN unreachable?) — geometry was "
                                          "measured with a fallback font; re-export with network access")
            return ex.run(Path(out_path))
        finally:
            browser.close()


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("html", help="finished AIJ poster.html")
    ap.add_argument("-o", "--output", help="output .pptx (default: <html stem>.pptx next to the HTML)")
    ap.add_argument("--report", help="also write the export report as JSON here")
    ap.add_argument("--mathjax-timeout-ms", type=int, default=15000)
    args = ap.parse_args(argv)

    html = Path(args.html)
    if not html.exists():
        print(f"error: {html} not found", file=sys.stderr)
        return 2
    try:
        import lxml  # noqa: F401
        import mathml2omml  # noqa: F401
        import pptx  # noqa: F401
    except ImportError as e:
        print(f"error: {e.name} is not installed — pip install python-pptx mathml2omml lxml "
              "(the skill's [pptx] extra)", file=sys.stderr)
        return 2
    out = Path(args.output) if args.output else html.with_suffix(".pptx")
    rep = export(html, out, mathjax_timeout_ms=args.mathjax_timeout_ms)
    print(f"[export_pptx] wrote {out}")
    print(f"  text frames: {rep.text_frames}  tables: {rep.tables}  pictures: {rep.pictures}  "
          f"equations: {rep.equations_native} display + {rep.equations_inline_native} inline native")
    for f in rep.equations_as_picture:
        print(f"  EQUATION AS PICTURE: {f['tex']!r} ({f['error']})")
    for f in rep.frames_as_picture:
        print(f"  FRAME AS PICTURE: {f['text']!r} ({f['reason']})")
    for w in rep.warnings:
        print(f"  WARN: {w}")
    print("  NOTE: fonts are referenced by name (SB Sans Display / SB Sans Display Light) and not embedded —"
          " install them on the machine that opens the .pptx.")
    if args.report:
        Path(args.report).write_text(json.dumps(rep.as_dict(), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
