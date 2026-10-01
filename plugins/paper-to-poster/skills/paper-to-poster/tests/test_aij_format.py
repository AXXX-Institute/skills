"""AIJ poster format: the scaffold, its frame fidelity, and its asset policy.

Static checks run everywhere. The render checks (computed type roles and frame
geometry vs. the organisers' .pptx) need Playwright/Chromium *and* network
access for the SB Sans Display web fonts (docs/adr/0011) — they skip otherwise.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import urllib.request
import zipfile
from pathlib import Path

import pytest

SKILL = Path(__file__).resolve().parent.parent
TEMPLATE_HTML = SKILL / "templates" / "portrait_aij.html"
TEMPLATE_PPTX = SKILL / "aij" / "assets" / "aij_template.pptx"
FONTS_CSS = SKILL / "aij" / "fonts.css"
EXAMPLE = SKILL / "examples" / "compression_horizon_aij" / "poster.html"
sys.path.insert(0, str(SKILL / "aij"))

EMU_PER_MM = 36000
PX_PER_MM = 96 / 25.4

#: Texts from the organisers' template that are instructions, not content.
INSTRUCTION_SNIPPETS = (
    "В шапке шаблона", "В нижней части", "Пространство по центру",
    "Просьба не менять", "Шрифт заголовка", "Шрифт основного текста",
    "Ромашка", "december@mail.ru", "Пестель",
)


# --------------------------------------------------------------------------
# static
# --------------------------------------------------------------------------

def test_template_canvas_is_the_organisers_slide():
    html = TEMPLATE_HTML.read_text()
    assert re.search(r"@page\s*{\s*size:\s*190\.5mm\s+275\.2mm", html)
    with zipfile.ZipFile(TEMPLATE_PPTX) as z:
        pres = z.read("ppt/presentation.xml").decode()
    m = re.search(r'<p:sldSz cx="(\d+)" cy="(\d+)"', pres)
    assert (int(m.group(1)) / EMU_PER_MM, round(int(m.group(2)) / EMU_PER_MM, 1)) == (190.5, 275.2)


def test_template_font_block_matches_fonts_css():
    import sync_fonts

    html = TEMPLATE_HTML.read_text()
    assert sync_fonts.sync(html) == html, "run: python aij/sync_fonts.py templates/portrait_aij.html"
    assert EXAMPLE.exists()
    assert sync_fonts.sync(EXAMPLE.read_text()) == EXAMPLE.read_text()


def test_fonts_come_from_the_cdn_only():
    css = FONTS_CSS.read_text()
    urls = re.findall(r'url\("([^"]+)"\)', css)
    assert len(urls) == 6
    assert all(u.startswith("https://cdn-app.sberdevices.ru/") for u in urls)
    assert {w for w in re.findall(r"font-weight:\s*(\d+)", css)} == {"300", "400", "600", "700"}


def test_no_font_binaries_in_the_skill():
    fonts = [p for p in SKILL.rglob("*")
             if p.suffix.lower() in (".woff", ".woff2", ".ttf", ".otf", ".eot", ".fntdata")]
    assert fonts == []


def test_bundled_template_has_its_embedded_fonts_stripped():
    with zipfile.ZipFile(TEMPLATE_PPTX) as z:
        names = z.namelist()
        pres = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()
        types = z.read("[Content_Types].xml").decode()
    assert not [n for n in names if n.startswith("ppt/fonts/")]
    assert "embeddedFont" not in pres and "embedTrueTypeFonts" not in pres
    assert "fonts/" not in rels and "fntdata" not in types
    # the frame survives byte-for-byte: background + AIJ mark
    assert (SKILL / "aij" / "assets" / "aij_background.png").read_bytes() == zipfile.ZipFile(TEMPLATE_PPTX).read("ppt/media/image1.png")
    assert (SKILL / "aij" / "assets" / "aij_mark.svg").read_bytes() == zipfile.ZipFile(TEMPLATE_PPTX).read("ppt/media/image6.svg")


def test_bundled_template_carries_no_sample_artwork():
    """The organisers' arrows, "COMPANY" sample-logo sheet and sample QR are
    blanked (the shapes stay, so the frame geometry is untouched)."""
    with zipfile.ZipFile(TEMPLATE_PPTX) as z:
        for part in ("ppt/media/image2.png", "ppt/media/image3.png", "ppt/media/image4.png",
                     "ppt/media/image5.png", "ppt/media/image7.svg"):
            assert z.getinfo(part).file_size < 200, part
        assert z.getinfo("docProps/thumbnail.jpeg").file_size < 20000


def test_template_has_no_pptx_instructions_and_sets_the_number():
    html = TEMPLATE_HTML.read_text()
    for s in INSTRUCTION_SNIPPETS:
        assert s not in html
    assert "№1" in html and "№TODO" not in html
    assert 'data-measure-role="poster"' in html and 'data-measure-role="footer"' in html


def test_prepare_assets_copies_the_frame(tmp_path):
    import prepare_assets

    out = prepare_assets.prepare(tmp_path / "images")
    assert sorted(p.name for p in out) == ["aij_background.png", "aij_mark.svg", "airi_5_years_logo_white.svg"]
    assert all(p.stat().st_size > 0 for p in out)


def test_sync_fonts_rewrites_a_stale_block(tmp_path):
    import sync_fonts

    stale = TEMPLATE_HTML.read_text().replace("SBSansDisplay-Light.woff2", "OLD-Light.woff2")
    p = tmp_path / "poster.html"
    p.write_text(stale)
    assert sync_fonts.main(["--check", str(p)]) == 1
    assert sync_fonts.main([str(p)]) == 0
    assert "OLD-Light" not in p.read_text()
    assert sync_fonts.main(["--check", str(p)]) == 0


# --------------------------------------------------------------------------
# rendered
# --------------------------------------------------------------------------

def _chromium_available() -> bool:
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as p:
            p.chromium.launch().close()
        return True
    except Exception:
        return False


def _cdn_reachable() -> bool:
    url = re.search(r'url\("([^"]+)"\)', FONTS_CSS.read_text()).group(1)
    try:
        with urllib.request.urlopen(urllib.request.Request(url, method="HEAD"), timeout=8) as r:
            return r.status == 200
    except Exception:
        return False


needs_render = pytest.mark.skipif(
    not (_chromium_available() and _cdn_reachable()),
    reason="needs playwright+chromium and network access to the SB Sans Display CDN",
)


@pytest.fixture(scope="module")
def scaffold_dir(tmp_path_factory):
    import prepare_assets

    d = tmp_path_factory.mktemp("aij_scaffold")
    shutil.copyfile(TEMPLATE_HTML, d / "poster.html")
    prepare_assets.prepare(d / "images")
    return d


def _render(html: Path, js: str):
    sys.path.insert(0, str(SKILL / "tools"))
    from _posterly import canvas as _canvas, render as _r
    from playwright.sync_api import sync_playwright

    _, viewport = _canvas.resolve_canvas(html, None, label="[test]")
    with sync_playwright() as p:
        browser, _ctx, page = _r.open_print_emulated_page(p, viewport)
        try:
            page.goto(html.as_uri())
            page.wait_for_load_state("networkidle")
            _r.settle_page(page)
            return page.evaluate(js)
        finally:
            browser.close()


_STYLE_JS = """
() => {
  const pick = (sel) => { const el = document.querySelector(sel); const s = getComputedStyle(el);
    return {family: s.fontFamily, weight: s.fontWeight, size: parseFloat(s.fontSize)}; };
  return {
    fonts: ['300','400','700'].map(w => document.fonts.check(w + ' 10px "SB Sans Display"')),
    title: pick('.aij-title'), number: pick('.aij-number'), authors: pick('.aij-authors'),
    heading: pick('.section-title'), body: pick('.section p'), li: pick('.section li'),
    caption: pick('.caption'), th: pick('.result-table th'), td: pick('.result-table td'),
    affiliations: pick('.aij-affiliations'), contact: pick('.aij-contact'),
  };
}
"""

PT = 96 / 72  # px per pt


@needs_render
def test_type_roles_render_in_sb_sans_display(scaffold_dir):
    got = _render(scaffold_dir / "poster.html", _STYLE_JS)
    assert got["fonts"] == [True, True, True], "SB Sans Display Light/Regular/Bold did not load"
    roles = {  # element -> (weight, pt)
        "title": ("600", 13), "number": ("600", 13),   # Semibold 13pt, as the template
        "authors": ("400", 7), "heading": ("700", 14), "th": ("400", 7),
        "body": ("400", 7), "li": ("400", 7), "caption": ("400", 7), "td": ("400", 7),
        "affiliations": ("300", 7), "contact": ("300", 7),
    }
    for key, (weight, pt) in roles.items():
        el = got[key]
        family = "SB Sans Text" if key in ("body", "li", "caption", "td", "th") else "SB Sans Display"
        assert el["family"].startswith(f'"{family}"'), (key, el)
        assert el["weight"] == weight, (key, el)
        assert el["size"] == pytest.approx(pt * PT, abs=0.05), (key, el)


def _pptx_frame() -> dict[str, tuple[float, float, float, float]]:
    """name -> (x, y, w, h) in mm, from the organisers' template; group
    children resolved to slide coordinates."""
    from pptx import Presentation

    prs = Presentation(str(TEMPLATE_PPTX))
    out = {}

    def walk(shapes, dx=0, dy=0):
        for s in shapes:
            if s.shape_type == 6:  # group
                xfrm = s._element.grpSpPr.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
                off, choff = xfrm[0], xfrm[2]
                walk(s.shapes, dx + int(off.get("x")) - int(choff.get("x")),
                     dy + int(off.get("y")) - int(choff.get("y")))
            else:
                out[s.name] = tuple(v / EMU_PER_MM for v in (s.left + dx, s.top + dy, s.width, s.height))
    walk(prs.slides[0].shapes)
    return out


_FRAME_JS = """
() => {
  const P = document.querySelector('.poster').getBoundingClientRect();
  const r = (sel) => { const el = document.querySelector(sel); if (!el) return null;
    const b = el.getBoundingClientRect(); return [b.x - P.x, b.y - P.y, b.width, b.height]; };
  return {bg: r('.aij-bg'), mark: r('.aij-mark'), number: r('.aij-number'), title: r('.aij-title'),
          authors: r('.aij-authors'), affiliations: r('.aij-affiliations'), contact: r('.aij-contact'),
          qr: r('.aij-qr'), logos: r('.aij-logos')};
}
"""

INSET_MM = 1.27  # PowerPoint's default top inset for the template's text boxes


@needs_render
@pytest.mark.parametrize("which", ["scaffold", "example"])
def test_frame_sits_on_the_pptx_coordinates(which, scaffold_dir):
    html = scaffold_dir / "poster.html" if which == "scaffold" else EXAMPLE
    got = {k: (tuple(v / PX_PER_MM for v in b) if b else None) for k, b in _render(html, _FRAME_JS).items()}
    ref = _pptx_frame()
    tol = 1.0

    def near(a, b):
        return all(abs(x - y) <= tol for x, y in zip(a, b))

    assert near(got["bg"], ref["Рисунок 10"]), got["bg"]
    assert near(got["mark"], ref["Рисунок 8"]), got["mark"]
    for key, name in (("number", "TextBox 70"), ("title", "TextBox 2"), ("authors", "TextBox 9")):
        x, y, _w, _h = ref[name]
        assert near(got[key][:2], (x, y + INSET_MM)), (key, got[key][:2], (x, y + INSET_MM))
    lx, ly = ref["Рисунок 13"][:2]            # first sample-logo slot: the logo band's corner
    assert abs(got["logos"][0] - lx) < tol, got["logos"]
    # Footer vertical positions are intentionally centred with the QR, not top-aligned.
    if which == "example":                    # one QR -> the outer tile
        assert near(got["qr"], ref["Скругленный прямоугольник 66"]), got["qr"]


@needs_render
def test_example_pdf_embeds_sb_sans_display():
    pdf = EXAMPLE.with_name("poster.pdf")
    assert pdf.exists()
    if not shutil.which("pdffonts"):
        pytest.skip("pdffonts (poppler-utils) not installed")
    fonts = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True, check=True).stdout
    for face in ("SBSansDisplay-Light", "SBSansDisplay-Regular", "SBSansDisplay-Semibold", "SBSansDisplay-Bold", "SBSansText-Regular", "SBSansText-Semibold"):
        assert face in fonts, fonts


@needs_render
def test_scaffold_passes_preflight(scaffold_dir):
    r = subprocess.run([sys.executable, str(SKILL / "tools" / "poster_check.py"), "preflight",
                        str(scaffold_dir / "poster.html")], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr


@needs_render
@pytest.mark.parametrize("gate", ["preflight", "measure", "polish"])
def test_example_passes_gates(gate):
    args = [sys.executable, str(SKILL / "tools" / "poster_check.py"), gate, str(EXAMPLE)]
    if gate == "polish":
        args.append("--strict")
    r = subprocess.run(args, capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr


@needs_render
def test_example_pdf_passes_verify_final():
    r = subprocess.run([sys.executable, str(SKILL / "tools" / "poster_check.py"), "verify-final",
                        str(EXAMPLE.with_name("poster.pdf")), "--from-html", str(EXAMPLE)],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr



# --------------------------------------------------------------------------
# white affiliation logos (the template's footer logos are white, no plate)
# --------------------------------------------------------------------------

_TWO_TONE = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50">
<style>.st0{fill:#002F87;} .st1{fill:#FFFFFF;}</style>
<circle class="st0" cx="25" cy="25" r="25"/><path class="st1" d="M20 10 H30 V40 H20 Z"/>
<path d="M60 10 H90 V40 H60 Z"/></svg>"""


SVG_NS = {"s": "http://www.w3.org/2000/svg"}


def _whiten(svg):
    pytest.importorskip("lxml")
    import white_logo
    from lxml import etree

    return etree.fromstring(white_logo.whiten_svg(svg).encode())


def test_white_logo_turns_colours_white_and_white_details_into_cut_outs():
    out = _whiten(_TWO_TONE)
    mask = out.find(".//s:mask", SVG_NS)
    rect = out.find("s:rect", SVG_NS)
    assert rect.get("mask") == f"url(#{mask.get('id')})"
    assert "fill:#FFFFFF !important" in rect.get("style")      # beats any rect{fill:…} in the logo's CSS
    css = mask.find(".//s:style", SVG_NS).text.replace(" ", "")
    assert ".st0{fill:#FFFFFF;}" in css                        # blue plate -> shown (white)
    assert ".st1{fill:#000000;}" in css                        # white letter -> cut out
    assert mask.find("s:g", SVG_NS).get("fill") == "#FFFFFF"   # unpainted (black) shapes -> white
    assert out.get("data-aij-white") == "1"


def test_white_logo_leaves_an_all_white_logo_alone():
    pytest.importorskip("lxml")
    import white_logo

    svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><path fill="#fff" d="M0 0H10V10Z"/></svg>'
    assert white_logo.whiten_svg(svg) == svg


def test_white_logo_is_idempotent():
    pytest.importorskip("lxml")
    import white_logo

    once = white_logo.whiten_svg(_TWO_TONE)
    assert white_logo.whiten_svg(once) == once


def test_white_logo_keeps_existing_masks_and_outline_artwork():
    """Figma exports: an inside-stroke <mask fill="white"> must stay a mask
    (not be inverted), and a root fill="none" (outline artwork) must not turn
    into solid shapes."""
    figma = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 40 40" fill="none">'
             '<mask id="m1" fill="white"><path d="M2 2H38V38H2Z"/></mask>'
             '<path d="M2 2H38V38H2Z" stroke="#123456" stroke-width="8" mask="url(#m1)"/></svg>')
    out = _whiten(figma)
    assert out.find(".//s:mask[@id='m1']", SVG_NS).get("fill") == "white"
    assert out.get("fill") == "none"
    art = [m for m in out.iter("{http://www.w3.org/2000/svg}mask") if m.get("id") != "m1"][0].find("s:g", SVG_NS)
    assert art.get("fill") is None                             # inherits the root's fill="none"


def test_white_logo_recognises_every_spelling_of_white():
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 90 30"><rect width="90" height="30" fill="#00A"/>'
           '<circle r="8" style="fill:rgba(255,255,255,1) !important"/><circle r="8" fill="#ffffffff"/>'
           '<circle r="8" fill="hsl(0, 0%, 100%)"/><circle r="8" fill="#FFF"/></svg>')
    circles = list(_whiten(svg).iter("{http://www.w3.org/2000/svg}circle"))
    assert "fill:#000000 !important" in circles[0].get("style")
    assert [c.get("fill") for c in circles[1:]] == ["#000000"] * 3


def test_white_logo_without_viewbox_and_with_embedded_raster():
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="20mm" '
           'height="10mm"><circle cx="38" cy="19" r="15" fill="#123"/><image href="data:image/png;base64,AA==" '
           'width="10" height="10"/></svg>')
    out = _whiten(svg)
    rect = out.find("s:rect", SVG_NS)
    assert (rect.get("width"), rect.get("height")) == ("120%", "120%")   # overshoots the viewport
    img = next(out.iter("{http://www.w3.org/2000/svg}image"))
    assert img.get("filter", "").startswith("url(#aij-white-image")


def test_white_logo_raster_keeps_alpha(tmp_path):
    pytest.importorskip("PIL")
    from PIL import Image
    import white_logo

    im = Image.new("RGBA", (4, 4), (0, 47, 135, 0))
    im.putpixel((1, 1), (0, 47, 135, 255))
    src = tmp_path / "logo.png"
    im.save(src)
    dst = tmp_path / "white.png"
    white_logo.whiten_raster(src, dst)  # test colour/alpha conversion before margin trimming
    out = Image.open(dst).convert("RGBA")
    assert out.getpixel((1, 1)) == (255, 255, 255, 255) and out.getpixel((0, 0))[3] == 0


def test_white_logo_opaque_raster_keeps_light_marks_solid(tmp_path):
    pytest.importorskip("PIL")
    from PIL import Image, ImageDraw
    import white_logo

    im = Image.new("RGB", (60, 30), "white")
    ImageDraw.Draw(im).rectangle((10, 5, 50, 25), fill=(255, 210, 0))   # yellow mark on white
    src = tmp_path / "yellow.png"
    im.save(src)
    dst = tmp_path / "white.png"
    white_logo.whiten_raster(src, dst)  # test colour/alpha conversion before margin trimming
    out = Image.open(dst).convert("RGBA")
    assert out.getpixel((30, 15)) == (255, 255, 255, 255) and out.getpixel((2, 2))[3] == 0


def test_white_logo_mask_overshoots_the_viewbox():
    """The soft-mask edge must lie outside the visible logo (PDF viewers can draw
    a seam along it)."""
    out = _whiten(_TWO_TONE)                       # viewBox 0 0 100 50
    mask = out.find(".//s:mask", SVG_NS)
    assert [float(mask.get(k)) for k in ("x", "y", "width", "height")] == [-10, -5, 120, 60]


def test_white_logo_renders_svg_to_a_high_res_png(tmp_path):
    pytest.importorskip("lxml")
    if not _chromium_available():
        pytest.skip("needs playwright + chromium")
    from PIL import Image
    import white_logo

    src = tmp_path / "logo.svg"
    src.write_text(_TWO_TONE)
    out = white_logo.make_white(src)
    assert out.name == "logo_white.png"
    im = Image.open(out)
    assert im.mode == "RGBA" and im.height == white_logo.PNG_HEIGHT_PX
    assert im.getchannel("A").getextrema() == (0, 255)
    assert white_logo.make_white(src, svg=True).name == "logo_white.svg"


def test_example_uses_white_logos_without_plates():
    html = EXAMPLE.read_text()
    logos = re.findall(r'<div class="aij-logo"><img src="images/([^"]+)"', html)
    assert logos[0] == "airi_5_years_logo_white.svg"
    assert all(l.endswith("_white.png") for l in logos[1:])
    assert not any("fusionbrain" in l or "axxx" in l for l in logos)
    assert all((EXAMPLE.parent / "images" / l).exists() for l in logos)
    assert "aij-logo-chip" not in html and "aij-logo-chip" not in TEMPLATE_HTML.read_text()



_LOGOS_JS = """
() => {
  const band = document.querySelector('.aij-logos').getBoundingClientRect();
  const imgs = [...document.querySelectorAll('.aij-logos img')].map(i => i.getBoundingClientRect());
  return {band: [band.x, band.y, band.width, band.height], imgs: imgs.map(r => [r.x, r.y, r.width, r.height])};
}
"""


@needs_render
def test_footer_logos_share_one_height_and_fit_the_band():
    """Designer review: logos are scaled to ONE common height (their widths follow
    their own proportions), not each fitted into an equal cell."""
    got = _render(EXAMPLE, _LOGOS_JS)
    band = [v / PX_PER_MM for v in got["band"]]
    imgs = [[v / PX_PER_MM for v in r] for r in got["imgs"]]
    assert len(imgs) == 3
    heights = [r[3] for r in imgs]
    assert max(heights) - min(heights) < 0.05, heights
    assert max(heights) <= 10.853 + 0.01
    assert imgs[-1][0] + imgs[-1][2] <= band[0] + band[2] + 0.1     # the row fits the band
    centres = [r[1] + r[3] / 2 for r in imgs]
    assert abs((imgs[0][0] + imgs[-1][0] + imgs[-1][2]) / 2 - (band[0] + band[2] / 2)) < 0.05
    gaps = [b[0] - a[0] - a[2] for a, b in zip(imgs, imgs[1:])]
    assert max(gaps) - min(gaps) < 0.05
    assert max(centres) - min(centres) < 0.05                          # and sits on one line


_HEADINGS_JS = """
() => ['.section-title', '.section-title .num', '.result-table th']
  .map(s => getComputedStyle(document.querySelector(s)).color)
"""


@needs_render
@pytest.mark.parametrize("which", ["scaffold", "example"])
def test_headings_are_black(which, scaffold_dir):
    html = scaffold_dir / "poster.html" if which == "scaffold" else EXAMPLE
    assert _render(html, _HEADINGS_JS) == ["rgb(0, 0, 0)"] * 3


@needs_render
def test_demo_content_inside_original_guides_and_footer_centred():
    from lxml import etree
    with zipfile.ZipFile(TEMPLATE_PPTX) as z:
        xml = etree.fromstring(z.read("ppt/viewProps.xml"))
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    guides = xml.findall(".//p:guide", ns)
    # Author-specified content guides, from the original file, in 1/8 point units.
    vertical = sorted(int(g.get("pos", 0)) for g in guides if g.get("orient") != "horz")
    horizontal = sorted(int(g.get("pos", 0)) for g in guides if g.get("orient") == "horz")
    left, right = vertical[1:3]
    top, bottom = horizontal[2:4]
    got = _render(EXAMPLE, """() => {
      const P = document.querySelector('.poster').getBoundingClientRect();
      const box = el => { const r=el.getBoundingClientRect(); return [r.left-P.left,r.top-P.top,r.right-P.left,r.bottom-P.top]; };
      const footer = ['.aij-footer-text','.aij-logos','.aij-qrs'].map(s=>box(document.querySelector(s)));
      const content = [...document.querySelectorAll('.body-grid .section, .body-grid img, .body-grid table')].map(box);
      const sub = getComputedStyle(document.querySelector('.section-subtitle'));
      return {content,footer,sub:[sub.fontFamily,sub.fontSize,sub.fontWeight]};
    }""")
    for x1, y1, x2, y2 in got['content']:
        assert x1 >= left / 6 - .1 and x2 <= right / 6 + .1
        assert y1 >= top / 6 - .1 and y2 <= bottom / 6 + .1
    centres = [(b[1]+b[3])/2 for b in got['footer']]
    assert max(centres)-min(centres) < .2
    assert got['footer'][0][2] < got['footer'][1][0]
    assert got['footer'][1][2] < got['footer'][2][0]
    assert got['sub'][0].startswith('"SB Sans Display"')
    assert float(got['sub'][1][:-2]) == pytest.approx(10*PT,abs=.05)
    assert got['sub'][2] == '600'


def test_white_logo_trim_removes_outer_alpha_margin(tmp_path):
    from PIL import Image
    import white_logo
    p = tmp_path / 'partner.png'
    im = Image.new('RGBA',(100,50)); im.paste((255,255,255,255),(20,10,80,40)); im.save(p)
    white_logo.trim_png(p)
    with Image.open(p) as got:
        assert got.size == (1200,600)
        assert got.getchannel('A').getextrema() == (255,255)
